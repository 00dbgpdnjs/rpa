from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3) # data is B2:C11
# bar_chart = BarChart() # Set the chart type (Bar, Line, Pie, ..)
# bar_chart.add_data(bar_value)

# ws.add_chart(bar_chart, "E1") # (what, where)

# < Create anotehr chart in detail >
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3) # B1:C11
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # 계열1, 계열2 -> value of 1st row [영어, 수학]
line_chart.title = "성적표"
line_chart.style = 10
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"

ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")